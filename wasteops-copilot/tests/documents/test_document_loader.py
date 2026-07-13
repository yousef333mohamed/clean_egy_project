"""Format-specific text extraction tests."""

from pathlib import Path
from types import SimpleNamespace

import pytest
from docx import Document
from pypdf import PdfWriter

from app.ingestion import document_loader
from app.ingestion.document_loader import DocumentLoadError, load_document, load_pdf


def test_txt_utf8_bom_and_arabic_are_preserved(tmp_path: Path) -> None:
    path = tmp_path / "policy.txt"
    path.write_text("إجراء السلامة\n\nTruck TRK-01", encoding="utf-8-sig")
    assert "إجراء السلامة" in load_document(path).text


def test_markdown_headings_and_code_are_preserved(tmp_path: Path) -> None:
    path = tmp_path / "manual.md"
    path.write_text("# Safety\n\nUse BIN-01.\n\n```text\nCODE-7\n```", encoding="utf-8")
    loaded = load_document(path)
    assert loaded.blocks[0].section_title == "Safety"
    assert "CODE-7" in loaded.text


def test_docx_headings_and_tables(tmp_path: Path) -> None:
    path = tmp_path / "manual.docx"
    document = Document()
    document.add_heading("Inspection", level=1)
    document.add_paragraph("Check TRK-01")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Item"
    table.cell(0, 1).text = "Status"
    document.save(path)
    loaded = load_document(path)
    assert any(block.section_title == "Inspection" for block in loaded.blocks)
    assert "Item | Status" in loaded.text


def test_empty_and_encrypted_pdf_fail_safely(tmp_path: Path) -> None:
    empty = tmp_path / "empty.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=72, height=72)
    with empty.open("wb") as target:
        writer.write(target)
    with pytest.raises(DocumentLoadError, match="OCR may be required"):
        load_pdf(empty)

    encrypted = tmp_path / "encrypted.pdf"
    writer.encrypt("secret")
    with encrypted.open("wb") as target:
        writer.write(target)
    with pytest.raises(DocumentLoadError, match="Encrypted"):
        load_pdf(encrypted)


def test_pdf_page_boundaries_and_empty_page_warning(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "fake.pdf"
    path.write_bytes(b"fake")
    pages = [SimpleNamespace(extract_text=lambda: "Page one"), SimpleNamespace(extract_text=lambda: "")]
    reader = SimpleNamespace(is_encrypted=False, pages=pages, metadata={"/Title": "Demo"})
    monkeypatch.setattr(document_loader, "PdfReader", lambda _: reader)
    loaded = load_pdf(path)
    assert loaded.blocks[0].page_number == 1
    assert loaded.total_pages == 2
    assert loaded.warnings == ["Page 2 contained no extractable text."]
